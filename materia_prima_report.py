import sys
sys.stdout.reconfigure(encoding='utf-8')
import io
import datetime
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

from config import get_corte_doblez_dir, normalize_po
from db_manager import get_all_pos
from corte_doblez_sync import load_corte_doblez_databases

STANDARD_COLUMNS = [
    'CAL 10 GALV',
    'CAL 12 GALV',
    'CAL 14 GALV',
    'CAL 16 GALV',
    'CAL 10 DECAPADO',
    'CAL 12 DECAPADO',
    'CAL 14 DECAPADO',
    'CAL 16 DECAPADO'
]

KNOWN_DECAPADO_POS = [
    '2602-0711', '2603-2561', '2603-2809', '2603-2811', '2603-2815', 
    '2603-2836', '2603-2838', '2608-3186', '2608-3261'
]

def classify_material_and_calibre(of_name, of_desc="", cal_field="", po_val="", proy_val="", piezas_text=""):
    """
    Regla Oficial de Planta Sigrama:
    - Las OF de material GALVANIZADO llevan obligatoriamente 'GALV' en el nombre de la OF.
    - Las OF de material INOXIDABLE llevan 'INOX'.
    - Las OF de material ALUMINIO llevan 'ALUM'.
    - Todas las demás OF (que no llevan GALV) son DECAPADO (lámina rolada en frío para pintura electrostática ANSI 61).
    """
    of_upper = f"{of_name} {of_desc}".upper()
    comb = f"{of_upper} {cal_field} {po_val} {proy_val}".upper()
    full_text = f"{comb} {piezas_text}".upper()
    
    # 1. Calibre
    cal = None
    if re.search(r'\b(CAL\.?\s*10|10\s*GA|CAL10|10GACR)\b', full_text):
        cal = 'CAL 10'
    elif re.search(r'\b(CAL\.?\s*12|12\s*GA|CAL12|12GACR)\b', full_text):
        cal = 'CAL 12'
    elif re.search(r'\b(CAL\.?\s*14|14\s*GA|CAL14|14GACR)\b', full_text):
        cal = 'CAL 14'
    elif re.search(r'\b(CAL\.?\s*16|16\s*GA|CAL16|16GACR)\b', full_text):
        cal = 'CAL 16'
    elif re.search(r'\b(CAL\.?\s*18|18\s*GA|CAL18|18GACR)\b', full_text):
        cal = 'CAL 18'
    elif re.search(r'\b(CAL\.?\s*20|20\s*GA|CAL20|20GACR)\b', full_text):
        cal = 'CAL 20'
        
    # 2. Material (Regla Directa de Planta)
    if re.search(r'\b(INOX|INOXIDABLE|SS304|SS316)\b', of_upper):
        mat = 'INOX'
    elif re.search(r'\b(ALUM|ALUMINIO|AL5052)\b', of_upper):
        mat = 'ALUMINIO'
    elif 'GALV' in of_upper:
        # Cualquier OF con 'GALV' en el nombre es Galvanizado
        mat = 'GALV'
    else:
        # Todas las demás OF que no llevan 'GALV' son material Decapado / ANSI 61
        mat = 'DECAPADO'
        
    return mat, cal

def extract_orden_interna(of_n, po_raw, proy_int, proy_cli, of_d, po_map):
    """
    Determina con precisión el número de Orden Interna (INT-001 en adelante)
    y su valor numérico entero para ordenamiento natural ascendente.
    """
    comb = f"{of_n} {po_raw} {proy_int} {proy_cli} {of_d}".upper()
    
    # Caso especial conocido INT-0047 (PO 2608-3425 CLOUD / GALVANIZADO)
    if (
        '047' in po_raw.upper() or '047' in proy_int.upper() or '047' in proy_cli.upper() or
        'PPAP GALV' in po_raw.upper() or
        any(x in of_n for x in ['00067', '00068', '00069', '00070', '00071', '00072', '00073', '00082', '00084', '00085', '00087', '00088', '00090', '00091']) or
        '3425' in po_raw or '3425' in of_n
    ):
        return 'INT-0047', 47
        
    p_norm = normalize_po(po_raw)
    if p_norm in po_map:
        info = po_map[p_norm]
        id_str = str(info.get('id_interno', '')).strip()
        m = re.search(r'\d+', id_str)
        if m:
            return id_str, int(m.group())
            
    for p_k, info in po_map.items():
        if len(p_k) >= 6 and p_k in normalize_po(comb):
            id_str = str(info.get('id_interno', '')).strip()
            m = re.search(r'\d+', id_str)
            if m:
                return id_str, int(m.group())
                
    # Limpiar posibles calibres antes de buscar números para no confundir CAL 10, CAL 12, etc.
    clean_comb = re.sub(r'\bCAL\.?\s*\d+\b', '', comb)
    clean_po = re.sub(r'\bCAL\.?\s*\d+\b', '', po_raw.upper())
    
    # 1. Buscar en campo Proyecto o PO expresiones directas como "PO 047", "PO 050", "PO 003", "INT 047"
    m_po_match = re.search(r'\b(?:PO|INT|OC)[-\s]*0*(\d{1,4})\b', f"{proy_int} {proy_cli} {clean_po}".upper())
    if m_po_match:
        val = int(m_po_match.group(1))
        if val < 2000:
            return f"INT-{val:04d}", val
            
    # 2. Buscar prefijos de proyecto como "003.- OC 2602-0711" o "026.- OC 2603-2561"
    m_proy = re.search(r'^\s*0*(\d{1,4})\s*\.-', proy_int)
    if m_proy:
        num = int(m_proy.group(1))
        return f"INT-{num:04d}", num
        
    # 3. Buscar "INT-XXXX" o "INT XXXX"
    m_int = re.search(r'\bINT[-\s]*0*(\d{1,4})\b', clean_comb)
    if m_int:
        num = int(m_int.group(1))
        return f"INT-{num:04d}", num
        
    # 4. Buscar "026 - OC" o similar
    m_fallback = re.search(r'\b0*(\d{1,3})\s*-\s*OC\b', clean_comb)
    if m_fallback:
        num = int(m_fallback.group(1))
        return f"INT-{num:04d}", num

    # 5. Si el nombre de la OF indica folio "OF 00003" o "0F 00003" (menor a 200)
    m_of_num = re.search(r'\b[0O]F[-\s]*0*(\d{1,4})\b', of_n.upper())
    if m_of_num:
        val_of = int(m_of_num.group(1))
        if 1 <= val_of <= 200:
            return f"INT-{val_of:04d}", val_of

    return "SIN INT", 99999

def extract_of_consecutive_digits(of_str):
    """Extrae el número consecutivo de la OF en formato 2 dígitos (ej. '01', '02', '83')."""
    m = re.search(r'\b[0O]\.?F\.?[-\s]*0*(\d{1,4})\b', str(of_str), re.I)
    if m:
        val = int(m.group(1))
        return f"{val:02d}"
    m2 = re.search(r'\b0*(\d{1,4})\b', str(of_str))
    if m2:
        val = int(m2.group(1))
        return f"{val:02d}"
    return str(of_str).strip()

def build_materia_prima_data():
    """Construye el dataset detallado y pivoteado del reporte de materia prima con ordenamiento INT-001 en adelante."""
    df_pos = get_all_pos()
    po_map = {}
    for _, r in df_pos.iterrows():
        p_str = str(r['po']).strip()
        p_norm = normalize_po(p_str)
        id_int = str(r.get('id_interno', '')).strip()
        proy = str(r.get('proyecto', '')).strip()
        info = {
            'po': p_str,
            'id_interno': id_int,
            'proyecto': proy
        }
        if p_norm:
            po_map[p_norm] = info
        m_num = re.search(r'\d+', id_int)
        if m_num:
            num = int(m_num.group())
            po_map[f"PO0{num}"] = info
            po_map[f"PO{num}"] = info
            po_map[f"INT00{num}"] = info
            po_map[f"INT0{num}"] = info
            po_map[f"INT{num}"] = info
            po_map[f"{num}"] = info

    dbs = load_corte_doblez_databases()
    df_ord = dbs[0]
    df_piez = dbs[1] if len(dbs) > 1 else pd.DataFrame()
    df_nid = dbs[4] if len(dbs) > 4 else pd.DataFrame()
    
    if df_ord.empty:
        return pd.DataFrame(), pd.DataFrame()

    # Mapa de textos de piezas por OF (para detección de gacr / ANSI 61 / Galv)
    piezas_map = {}
    if not df_piez.empty and 'of_number' in df_piez.columns:
        for of_num, grp in df_piez.groupby('of_number'):
            p_names = grp['nombre_pieza'].dropna().astype(str).tolist()
            piezas_map[of_num] = " ".join(p_names).upper()

    raw_records = []
    for _, r in df_ord.iterrows():
        of_n = str(r['of_number']).strip()
        of_d = str(r.get('descripcion_pronest') or r.get('descripcion') or '').strip()
        po_raw = str(r.get('po') or '').strip()
        cal_raw = str(r.get('calibre') or '').strip()
        proy_int = str(r.get('proyecto') or '').strip()
        proy_cli = str(r.get('proyecto_cliente') or '').strip()
        proy_raw = proy_cli if proy_cli and proy_cli not in ('nan', 'POR DEFINIR') else proy_int
        
        po_norm = normalize_po(po_raw)
        po_match = None
        
        # 1. Alias integral INT-0047 (PO 2608-3425 CLOUD / GALVANIZADO)
        is_int_47 = (
            '047' in po_raw.upper() or '047' in proy_int.upper() or '047' in proy_cli.upper() or
            'PPAP GALV' in po_raw.upper() or
            any(x in of_n for x in ['00067', '00068', '00069', '00070', '00071', '00072', '00073', '00082', '00084', '00085', '00087', '00088', '00090', '00091']) or
            '3425' in po_norm or '3425' in of_n
        )
        if is_int_47:
            po_match = po_map.get('26083425')
            if not po_match:
                po_match = {'po': '26083425', 'id_interno': 'INT-0047', 'proyecto': 'CLOUD'}
        elif po_norm in po_map:
            po_match = po_map[po_norm]
        else:
            # Primero verificar si el texto contiene una PO formal de 8 dígitos (ej. 2603-2815)
            m_po_pat = re.search(r'\b(260\d-?\d{4})\b', f"{po_raw} {of_n} {proy_int}")
            if m_po_pat and normalize_po(m_po_pat.group(1)) in po_map:
                po_match = po_map[normalize_po(m_po_pat.group(1))]
            else:
                comb_srch = f"{of_n} {po_raw}".upper()
                m_id = re.search(r'\b(?:PO|INT|OC)?\s*0*(\d{1,3})\b', comb_srch)
                if m_id and m_id.group(1) in po_map:
                    po_match = po_map[m_id.group(1)]
                else:
                    for p_k, p_info in po_map.items():
                        if len(p_k) >= 6 and p_k in normalize_po(comb_srch):
                            po_match = p_info
                            break
                        
        if po_match:
            po_disp = po_match['po']
            proy_disp = po_match['proyecto']
        else:
            # Extracción de formato estándar de PO (ej. 2602-0711, 2603-2561)
            m_po_pat = re.search(r'\b(260\d-?\d{4})\b', f"{po_raw} {of_n}")
            if m_po_pat:
                po_disp = m_po_pat.group(1)
            else:
                po_disp = po_raw if po_raw and po_raw != 'nan' else 'Sin PO'
            
            if proy_raw and proy_raw not in ('nan', 'POR DEFINIR'):
                proy_disp = proy_raw
            elif 'RENO' in of_n:
                proy_disp = 'RENO 6'
            elif 'LC8' in of_n:
                proy_disp = 'LC8 20K'
            else:
                proy_disp = 'General'
                
        # Estandarización de nombres de proyectos reconocidos
        if '0711' in po_disp or '0711' in of_n:
            proy_disp = 'RENO 6'
        elif '2561' in po_disp or '2561' in of_n:
            proy_disp = 'LC8 20K'
        elif '2809' in po_disp or '2809' in of_n:
            proy_disp = 'ALM SWBD META'
        elif '2811' in po_disp or '2811' in of_n:
            proy_disp = 'ALM SWBD SOUTH VALLEY'
        elif '2815' in po_disp or '2815' in of_n:
            proy_disp = 'SWBD RENO 4'
            
        # Extracción precisa de Orden Interna (INT-001 en adelante) y valor numérico
        id_disp, int_num = extract_orden_interna(of_n, po_disp, proy_int, proy_cli, of_d, po_map)
            
        m_nid = df_nid[df_nid['of_number'] == of_n] if not df_nid.empty else pd.DataFrame()
        hojas = float(m_nid['hojas'].sum()) if not m_nid.empty and 'hojas' in m_nid.columns else 0.0
        nidos_cnt = len(m_nid) if not m_nid.empty else 0
        
        piezas_txt = piezas_map.get(of_n, '')
        mat, cal = classify_material_and_calibre(of_n, of_d, cal_raw, po_raw, proy_raw, piezas_txt)
        mat_key = f"{cal} {mat}" if (cal and mat) else "OTRO"
        
        raw_records.append({
            'id_interno': id_disp,
            'int_num': int_num,
            'po': po_disp,
            'of_number': of_n,
            'proyecto': proy_disp,
            'material': mat,
            'calibre': cal or 'N/A',
            'material_key': mat_key,
            'hojas': hojas,
            'nidos_count': nidos_cnt
        })

    df_raw = pd.DataFrame(raw_records)
    if df_raw.empty:
        return pd.DataFrame(), pd.DataFrame()

    # 1. PIVOT POR OF (Detalle idéntico a la imagen con Orden Interna al frente)
    pivot_of = df_raw.pivot_table(
        index=['id_interno', 'int_num', 'po', 'of_number', 'proyecto'],
        columns='material_key',
        values='hojas',
        aggfunc='sum',
        fill_value=0.0
    ).reset_index()

    for col in STANDARD_COLUMNS:
        if col not in pivot_of.columns:
            pivot_of[col] = 0.0

    col_order_of = ['id_interno', 'int_num', 'po', 'of_number', 'proyecto'] + STANDARD_COLUMNS
    extra_cols = [c for c in pivot_of.columns if c not in col_order_of and c != 'TOTAL HOJAS']
    col_order_of.extend(extra_cols)
    
    pivot_of['TOTAL HOJAS'] = pivot_of[STANDARD_COLUMNS + extra_cols].sum(axis=1)
    col_order_of.append('TOTAL HOJAS')
    pivot_of = pivot_of[[c for c in col_order_of if c in pivot_of.columns]]
    
    # Ordenar por defecto por Orden Interna ascendente (INT-001 en adelante)
    pivot_of = pivot_of.sort_values(by=['int_num', 'po', 'of_number'], ascending=[True, True, True]).reset_index(drop=True)

    # 2. PIVOT CONSOLIDADO POR PO
    pivot_po = df_raw.pivot_table(
        index=['id_interno', 'int_num', 'po', 'proyecto'],
        columns='material_key',
        values='hojas',
        aggfunc='sum',
        fill_value=0.0
    ).reset_index()

    for col in STANDARD_COLUMNS:
        if col not in pivot_po.columns:
            pivot_po[col] = 0.0

    # Mapeo de OFs consecutivas por PO (ej. '01, 02, 03')
    po_ofs_map = {}
    for po_val, grp in df_raw.groupby('po'):
        digits_set = set()
        for of_n in grp['of_number']:
            d = extract_of_consecutive_digits(of_n)
            if d:
                digits_set.add(d)
        def _sort_k(x):
            try:
                return (0, int(x))
            except:
                return (1, str(x))
        sorted_digits = sorted(list(digits_set), key=_sort_k)
        po_ofs_map[po_val] = ", ".join(sorted_digits) if sorted_digits else "Sin Información"

    pivot_po['ofs_count'] = pivot_po['po'].map(po_ofs_map).fillna("Sin Información")
    pivot_po['informacion'] = "Con Información"

    # Incorporar todas las órdenes internas de po_tracker.db que aún no tengan registros de corte
    existing_pos_norm = set(normalize_po(p) for p in pivot_po['po'])
    missing_records = []
    
    for _, r in df_pos.iterrows():
        p_str = str(r['po']).strip()
        p_norm = normalize_po(p_str)
        if p_norm not in existing_pos_norm:
            id_int = str(r.get('id_interno', '')).strip()
            proy = str(r.get('proyecto', '')).strip()
            m_num = re.search(r'\d+', id_int)
            int_num = int(m_num.group()) if m_num else 999
            
            row_dict = {
                'id_interno': id_int,
                'int_num': int_num,
                'po': p_str,
                'proyecto': proy,
                'ofs_count': 'Sin Información',
                'informacion': 'Sin Información'
            }
            for col in STANDARD_COLUMNS:
                row_dict[col] = 0.0
            row_dict['TOTAL HOJAS'] = 0.0
            missing_records.append(row_dict)
            existing_pos_norm.add(p_norm)
            
    if missing_records:
        df_missing = pd.DataFrame(missing_records)
        pivot_po = pd.concat([pivot_po, df_missing], ignore_index=True)

    base_cols_po = ['id_interno', 'int_num', 'po', 'proyecto', 'ofs_count', 'informacion']
    extra_cols_po = [c for c in pivot_po.columns if c not in base_cols_po and c not in STANDARD_COLUMNS and c != 'TOTAL HOJAS']
    
    for col in STANDARD_COLUMNS + extra_cols_po:
        pivot_po[col] = pd.to_numeric(pivot_po[col], errors='coerce').fillna(0.0)
        
    pivot_po['TOTAL HOJAS'] = pivot_po[STANDARD_COLUMNS + extra_cols_po].sum(axis=1)

    col_order_po = base_cols_po + STANDARD_COLUMNS + extra_cols_po + ['TOTAL HOJAS']
    pivot_po = pivot_po[[c for c in col_order_po if c in pivot_po.columns]]
    
    # Ordenar por defecto por Orden Interna ascendente (INT-001 en adelante)
    pivot_po = pivot_po.sort_values(by=['int_num', 'po'], ascending=[True, True]).reset_index(drop=True)

    return pivot_of, pivot_po

def generate_materia_prima_excel(df_pivot_of, df_pivot_po=None):
    """Genera archivo Excel idéntico al formato corporativo y al de la imagen del usuario."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Uso_Materia_Prima_OF"
    ws.views.sheetView[0].showGridLines = True
    
    C_HEADER_BLUE = "2B6CB0" # Azul idéntico a la imagen
    C_SLATE_DARK  = "0F172A"
    C_WHITE       = "FFFFFF"
    C_ZEBRA       = "F8FAFC"
    C_TOTAL_BG    = "EDF2F7"
    C_TOTAL_TEXT  = "1A202C"
    
    # ── 1. Banner Superior
    ws.merge_cells("A1:M1")
    ws.merge_cells("A2:M2")
    c1 = ws["A1"]
    c1.value = "INDUSTRIA SIGRAMA S.A. DE C.V.  —  REPORTE DE USO DE MATERIA PRIMA (LÁMINAS POR OF / PO)"
    c1.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    c1.fill = PatternFill(start_color=C_SLATE_DARK, end_color=C_SLATE_DARK, fill_type="solid")
    c1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    tot_hojas_glob = float(df_pivot_of['TOTAL HOJAS'].sum()) if not df_pivot_of.empty and 'TOTAL HOJAS' in df_pivot_of.columns else 0.0
    c2 = ws["A2"]
    c2.value = f"Control Central de Suministro | Emisión: {now_str} | Total OFs: {len(df_pivot_of)} | Consumo Total Acumulado: {tot_hojas_glob:,.0f} Hojas / Láminas"
    c2.font = Font(name="Calibri", size=9.5, italic=True, color="94A3B8")
    c2.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    c2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8

    # ── 2. Encabezados de Columna (Fila 4) - Idénticos a la imagen con Orden Interna
    headers = [
        ("ORDEN INTERNA", "center", 16),
        ("PO", "center", 14),
        ("OF", "left", 34),
        ("PROYECTO", "left", 16),
        ("CAL 10 GALV", "right", 14),
        ("CAL 12 GALV", "right", 14),
        ("CAL 14 GALV", "right", 14),
        ("CAL 16 GALV", "right", 14),
        ("CAL 10 DECAPADO", "right", 15),
        ("CAL 12 DECAPADO", "right", 15),
        ("CAL 14 DECAPADO", "right", 15),
        ("CAL 16 DECAPADO", "right", 15),
        ("TOTAL HOJAS", "right", 15),
    ]
    
    hdr_row = 4
    ws.row_dimensions[hdr_row].height = 28
    border_hdr = Border(
        top=Side(border_style="thin", color="CBD5E1"),
        bottom=Side(border_style="medium", color="1A365D"),
        left=Side(border_style="thin", color="4299E1"),
        right=Side(border_style="thin", color="4299E1")
    )
    
    for col_i, (h_text, align_h, col_w) in enumerate(headers, start=1):
        cell = ws.cell(row=hdr_row, column=col_i)
        cell.value = h_text
        cell.font = Font(name="Calibri", size=10, bold=True, color=C_WHITE)
        cell.fill = PatternFill(start_color=C_HEADER_BLUE, end_color=C_HEADER_BLUE, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_hdr
        ws.column_dimensions[get_column_letter(col_i)].width = col_w

    # ── 3. Filas de Datos
    start_row = 5
    border_data = Border(
        top=Side(border_style="thin", color="E2E8F0"),
        bottom=Side(border_style="thin", color="E2E8F0"),
        left=Side(border_style="thin", color="E2E8F0"),
        right=Side(border_style="thin", color="E2E8F0")
    )
    
    for idx, (_, r) in enumerate(df_pivot_of.iterrows()):
        curr_row = start_row + idx
        ws.row_dimensions[curr_row].height = 20
        bg_color = C_ZEBRA if (idx % 2 == 1) else "FFFFFF"
        fill_row = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        
        id_v   = str(r.get('id_interno', '')).strip()
        po_v   = str(r.get('po', '')).strip()
        of_v   = str(r.get('of_number', '')).strip()
        proy_v = str(r.get('proyecto', '')).strip()
        
        vals = [
            (id_v,   "center", "@", Font(name="Calibri", size=9.5, bold=True, color="1E3A8A")),
            (po_v,   "center", "@", Font(name="Calibri", size=9.5, bold=True, color="EC2024")),
            (of_v,   "left",   "@", Font(name="Calibri", size=9, bold=True, color="1E293B")),
            (proy_v, "left",   "@", Font(name="Calibri", size=9.5, color="334155")),
        ]
        
        for col_name in STANDARD_COLUMNS:
            h_val = float(r.get(col_name, 0) or 0)
            if h_val > 0:
                fnt_h = Font(name="Calibri", size=9.5, bold=True, color="2B6CB0")
                vals.append((h_val, "right", '#,##0', fnt_h))
            else:
                fnt_h = Font(name="Calibri", size=9, color="CBD5E1")
                vals.append(("-", "center", "@", fnt_h))
                
        tot_h = float(r.get('TOTAL HOJAS', 0) or 0)
        vals.append((tot_h, "right", '#,##0 "hjs"', Font(name="Calibri", size=10, bold=True, color="0F172A")))
        
        for c_idx, (val, al, nf, fnt) in enumerate(vals, start=1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.value = val
            cell.alignment = Alignment(horizontal=al, vertical="center")
            cell.number_format = nf
            cell.font = fnt
            cell.fill = fill_row
            cell.border = border_data

    # ── 4. Fila de Totales
    end_row = start_row + len(df_pivot_of) - 1
    tot_row = end_row + 1
    ws.row_dimensions[tot_row].height = 24
    
    ws.merge_cells(f"A{tot_row}:D{tot_row}")
    c_tot = ws[f"A{tot_row}"]
    c_tot.value = "TOTALES GENERALES DE HOJAS"
    c_tot.font = Font(name="Calibri", size=10, bold=True, color=C_TOTAL_TEXT)
    c_tot.alignment = Alignment(horizontal="center", vertical="center")
    
    border_total = Border(
        top=Side(border_style="thin", color="0F172A"),
        bottom=Side(border_style="double", color="0F172A"),
        left=Side(border_style="thin", color="CBD5E1"),
        right=Side(border_style="thin", color="CBD5E1")
    )
    
    for c_idx in range(1, 5):
        cell = ws.cell(row=tot_row, column=c_idx)
        cell.fill = PatternFill(start_color=C_TOTAL_BG, end_color=C_TOTAL_BG, fill_type="solid")
        cell.border = border_total

    for c_idx in range(5, 14):
        col_letter = get_column_letter(c_idx)
        cell = ws.cell(row=tot_row, column=c_idx)
        cell.value = f"=SUM({col_letter}{start_row}:{col_letter}{end_row})"
        cell.number_format = '#,##0' if c_idx < 12 else '#,##0 "hjs"'
        cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        cell.fill = PatternFill(start_color=C_TOTAL_BG, end_color=C_TOTAL_BG, fill_type="solid")
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = border_total

    # ── 5. Hoja 2: Resumen Consolidado por PO
    if df_pivot_po is not None and not df_pivot_po.empty:
        ws2 = wb.create_sheet(title="Consolidado_por_PO")
        ws2.views.sheetView[0].showGridLines = True
        
        ws2.merge_cells("A1:N1")
        c2_1 = ws2["A1"]
        c2_1.value = "INDUSTRIA SIGRAMA S.A. DE C.V.  —  CONSUMO DE MATERIA PRIMA POR ORDEN DE COMPRA (PO)"
        c2_1.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        c2_1.fill = PatternFill(start_color=C_SLATE_DARK, end_color=C_SLATE_DARK, fill_type="solid")
        c2_1.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 24
        
        headers_po = [
            ("ID Interno", 12, "center"),
            ("PO / Folio", 14, "center"),
            ("Proyecto", 16, "left"),
            ("OFs #", 24, "center"),
            ("Información", 16, "center"),
            ("CAL 10 GALV", 14, "right"),
            ("CAL 12 GALV", 14, "right"),
            ("CAL 14 GALV", 14, "right"),
            ("CAL 16 GALV", 14, "right"),
            ("CAL 10 DECAPADO", 15, "right"),
            ("CAL 12 DECAPADO", 15, "right"),
            ("CAL 14 DECAPADO", 15, "right"),
            ("CAL 16 DECAPADO", 15, "right"),
            ("TOTAL HOJAS", 15, "right"),
        ]
        
        ws2.row_dimensions[2].height = 26
        for c_i, (h_t, col_w, al_h) in enumerate(headers_po, start=1):
            cell = ws2.cell(row=2, column=c_i)
            cell.value = h_t
            cell.font = Font(name="Calibri", size=10, bold=True, color=C_WHITE)
            cell.fill = PatternFill(start_color=C_HEADER_BLUE, end_color=C_HEADER_BLUE, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_hdr
            ws2.column_dimensions[get_column_letter(c_i)].width = col_w
            
        po_start = 3
        for idx_p, (_, r_p) in enumerate(df_pivot_po.iterrows()):
            curr_r = po_start + idx_p
            ws2.row_dimensions[curr_r].height = 20
            bg_p = C_ZEBRA if (idx_p % 2 == 1) else "FFFFFF"
            fill_p = PatternFill(start_color=bg_p, end_color=bg_p, fill_type="solid")
            
            p_id = str(r_p.get('id_interno', '')).strip()
            p_po = str(r_p.get('po', '')).strip()
            p_proy = str(r_p.get('proyecto', '')).strip()
            p_ofs = str(r_p.get('ofs_count', 'Sin Información')).strip()
            p_info = str(r_p.get('informacion', 'Sin Información')).strip()
            
            fnt_ofs = Font(name="Calibri", size=9, bold=(p_ofs != "Sin Información"), color="475569" if p_ofs != "Sin Información" else "94A3B8")
            fnt_info = Font(name="Calibri", size=9, bold=True, color="166534" if p_info == "Con Información" else "94A3B8")
            
            vals_p = [
                (p_id,   "center", "@", Font(name="Calibri", size=9.5, bold=True, color="0F172A")),
                (p_po,   "center", "@", Font(name="Calibri", size=9.5, bold=True, color="EC2024")),
                (p_proy, "left",   "@", Font(name="Calibri", size=9.5, color="334155")),
                (p_ofs,  "center", "@", fnt_ofs),
                (p_info, "center", "@", fnt_info),
            ]
            
            for col_name in STANDARD_COLUMNS:
                h_val = float(r_p.get(col_name, 0) or 0)
                if h_val > 0:
                    vals_p.append((h_val, "right", '#,##0', Font(name="Calibri", size=9.5, bold=True, color="2B6CB0")))
                else:
                    vals_p.append(("-", "center", "@", Font(name="Calibri", size=9, color="CBD5E1")))
                    
            tot_h_p = float(r_p.get('TOTAL HOJAS', 0) or 0)
            if tot_h_p > 0:
                vals_p.append((tot_h_p, "right", '#,##0 "hjs"', Font(name="Calibri", size=10, bold=True, color="0F172A")))
            else:
                vals_p.append(("-", "center", "@", Font(name="Calibri", size=9, color="CBD5E1")))
            
            for c_i, (val, al, nf, fnt) in enumerate(vals_p, start=1):
                cell = ws2.cell(row=curr_r, column=c_i)
                cell.value = val
                cell.alignment = Alignment(horizontal=al, vertical="center")
                cell.number_format = nf
                cell.font = fnt
                cell.fill = fill_p
                cell.border = border_data
                
        po_end = po_start + len(df_pivot_po) - 1
        po_tot = po_end + 1
        ws2.row_dimensions[po_tot].height = 24
        ws2.merge_cells(f"A{po_tot}:E{po_tot}")
        c_p_tot = ws2[f"A{po_tot}"]
        c_p_tot.value = "TOTALES CONSOLIDADOS POR PO"
        c_p_tot.font = Font(name="Calibri", size=10, bold=True, color=C_TOTAL_TEXT)
        c_p_tot.alignment = Alignment(horizontal="center", vertical="center")
        
        for c_i in range(1, 6):
            cell = ws2.cell(row=po_tot, column=c_i)
            cell.fill = PatternFill(start_color=C_TOTAL_BG, end_color=C_TOTAL_BG, fill_type="solid")
            cell.border = border_total
            
        for c_i in range(6, 15):
            col_l = get_column_letter(c_i)
            cell = ws2.cell(row=po_tot, column=c_i)
            cell.value = f"=SUM({col_l}{po_start}:{col_l}{po_end})"
            cell.number_format = '#,##0' if c_i < 14 else '#,##0 "hjs"'
            cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
            cell.fill = PatternFill(start_color=C_TOTAL_BG, end_color=C_TOTAL_BG, fill_type="solid")
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = border_total

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_materia_prima_page():
    """Renderiza la vista completa del Reporte de Uso de Materia Prima por PO / OF."""
    import streamlit as st
    import streamlit.components.v1 as components
    
    # ── 1. Encabezado Oficial
    col_h1, col_h2 = st.columns([3.3, 1.2])
    with col_h1:
        st.markdown("""
        <div style="background: #FFFFFF; border: 1px solid #CBD5E1; border-left: 6px solid #2B6CB0; border-radius: 10px; padding: 18px 24px; margin-bottom: 18px; box-shadow: 0 2px 8px rgba(0,0,0,0.06);">
            <h2 style="color: #0F172A !important; font-family: 'Montserrat', sans-serif; font-size: 22px; font-weight: 800; margin: 0 0 4px 0; letter-spacing: 0.5px;">
                🏭 Reporte de Uso de Materia Prima por PO (Láminas)
            </h2>
            <p style="color: #475569 !important; font-size: 13.5px; margin: 0; font-family: 'Questrial', sans-serif;">
                Consumo y programación de láminas en Taller de Corte (Pronest) desglosado por material (<b style="color:#2B6CB0;">Galvanizado</b>, <b style="color:#475569;">Decapado</b>, Inox) y calibre por Orden de Compra y Fabricación.
            </p>
        </div>
        """, unsafe_allow_html=True)
    with col_h2:
        st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)
        if st.button("🔄 Sincronizar en Vivo", type="primary", use_container_width=True, key="btn_sync_materia_prima", help="Actualiza avances de taller y nidos de corte en tiempo real"):
            with st.spinner("Sincronizando con Planta y Taller de Corte..."):
                from remisiones_sync import sync_live_remisiones_from_github
                sync_live_remisiones_from_github()
                st.cache_data.clear()
            st.toast("✅ ¡Datos de Materia Prima sincronizados con éxito!")
            st.rerun()

    df_of, df_po = build_materia_prima_data()
    if df_of.empty:
        st.info("💡 No se encontraron órdenes de fabricación en la base de datos de Corte.")
        return

    # ── 2. Tarjetas de Totales Globales (KPIs)
    tot_hjs_glob = float(df_of['TOTAL HOJAS'].sum())
    tot_cal10_g = float(df_of['CAL 10 GALV'].sum())
    tot_cal12_g = float(df_of['CAL 12 GALV'].sum())
    tot_cal14_g = float(df_of['CAL 14 GALV'].sum())
    tot_cal16_g = float(df_of['CAL 16 GALV'].sum())
    tot_decap_glob = float(df_of[['CAL 10 DECAPADO', 'CAL 12 DECAPADO', 'CAL 14 DECAPADO', 'CAL 16 DECAPADO']].sum().sum())
    tot_ofs_cnt = len(df_of)
    tot_pos_cnt = len(df_po)
    pos_con_info = len(df_po[df_po['informacion'] == 'Con Información']) if 'informacion' in df_po.columns else df_of['po'].nunique()

    k1, k2, k3, k4, k5, k6 = st.columns(6)
    with k1:
        st.markdown(f"""
        <div style="background:#FFFFFF; border:1px solid #E2E8F0; border-top:4px solid #0F172A; border-radius:10px; padding:12px 14px; box-shadow:0 3px 6px rgba(0,0,0,0.04); min-height:105px;">
            <div style="font-size:11px; font-weight:800; color:#0F172A; text-transform:uppercase;">Total Hojas</div>
            <div style="font-size:24px; font-weight:900; color:#0F172A; margin:4px 0 2px 0;">{tot_hjs_glob:,.0f} <span style="font-size:12px; font-weight:500; color:#64748B;">hjs</span></div>
            <span style="background:#F1F5F9; color:#475569; font-size:10px; font-weight:700; padding:2px 7px; border-radius:10px;">{tot_ofs_cnt} OFs | {pos_con_info} con Corte / {tot_pos_cnt} POs</span>
        </div>
        """, unsafe_allow_html=True)
    with k2:
        st.markdown(f"""
        <div style="background:#FFFFFF; border:1px solid #E2E8F0; border-top:4px solid #2B6CB0; border-radius:10px; padding:12px 14px; box-shadow:0 3px 6px rgba(0,0,0,0.04); min-height:105px;">
            <div style="font-size:11px; font-weight:800; color:#2B6CB0; text-transform:uppercase;">Cal. 10 Galv</div>
            <div style="font-size:24px; font-weight:900; color:#0F172A; margin:4px 0 2px 0;">{tot_cal10_g:,.0f} <span style="font-size:12px; font-weight:500; color:#64748B;">hjs</span></div>
            <span style="background:#EBF8FF; color:#2B6CB0; font-size:10px; font-weight:700; padding:2px 7px; border-radius:10px;">{(tot_cal10_g/tot_hjs_glob*100):.1f}% del total</span>
        </div>
        """, unsafe_allow_html=True)
    with k3:
        st.markdown(f"""
        <div style="background:#FFFFFF; border:1px solid #E2E8F0; border-top:4px solid #3182CE; border-radius:10px; padding:12px 14px; box-shadow:0 3px 6px rgba(0,0,0,0.04); min-height:105px;">
            <div style="font-size:11px; font-weight:800; color:#3182CE; text-transform:uppercase;">Cal. 12 Galv</div>
            <div style="font-size:24px; font-weight:900; color:#0F172A; margin:4px 0 2px 0;">{tot_cal12_g:,.0f} <span style="font-size:12px; font-weight:500; color:#64748B;">hjs</span></div>
            <span style="background:#EBF8FF; color:#3182CE; font-size:10px; font-weight:700; padding:2px 7px; border-radius:10px;">{(tot_cal12_g/tot_hjs_glob*100):.1f}% del total</span>
        </div>
        """, unsafe_allow_html=True)
    with k4:
        st.markdown(f"""
        <div style="background:#FFFFFF; border:1px solid #E2E8F0; border-top:4px solid #4299E1; border-radius:10px; padding:12px 14px; box-shadow:0 3px 6px rgba(0,0,0,0.04); min-height:105px;">
            <div style="font-size:11px; font-weight:800; color:#4299E1; text-transform:uppercase;">Cal. 14 Galv</div>
            <div style="font-size:24px; font-weight:900; color:#0F172A; margin:4px 0 2px 0;">{tot_cal14_g:,.0f} <span style="font-size:12px; font-weight:500; color:#64748B;">hjs</span></div>
            <span style="background:#EBF8FF; color:#4299E1; font-size:10px; font-weight:700; padding:2px 7px; border-radius:10px;">{(tot_cal14_g/tot_hjs_glob*100):.1f}% del total</span>
        </div>
        """, unsafe_allow_html=True)
    with k5:
        st.markdown(f"""
        <div style="background:#FFFFFF; border:1px solid #E2E8F0; border-top:4px solid #63B3ED; border-radius:10px; padding:12px 14px; box-shadow:0 3px 6px rgba(0,0,0,0.04); min-height:105px;">
            <div style="font-size:11px; font-weight:800; color:#2B6CB0; text-transform:uppercase;">Cal. 16 Galv</div>
            <div style="font-size:24px; font-weight:900; color:#0F172A; margin:4px 0 2px 0;">{tot_cal16_g:,.0f} <span style="font-size:12px; font-weight:500; color:#64748B;">hjs</span></div>
            <span style="background:#EBF8FF; color:#2B6CB0; font-size:10px; font-weight:700; padding:2px 7px; border-radius:10px;">{(tot_cal16_g/tot_hjs_glob*100):.1f}% del total</span>
        </div>
        """, unsafe_allow_html=True)
    with k6:
        st.markdown(f"""
        <div style="background:#FFFFFF; border:1px solid #E2E8F0; border-top:4px solid #475569; border-radius:10px; padding:12px 14px; box-shadow:0 3px 6px rgba(0,0,0,0.04); min-height:105px;">
            <div style="font-size:11px; font-weight:800; color:#475569; text-transform:uppercase;">Decapado / ANSI 61</div>
            <div style="font-size:24px; font-weight:900; color:#0F172A; margin:4px 0 2px 0;">{tot_decap_glob:,.0f} <span style="font-size:12px; font-weight:500; color:#64748B;">hjs</span></div>
            <span style="background:#F1F5F9; color:#475569; font-size:10px; font-weight:700; padding:2px 7px; border-radius:10px;">{(tot_decap_glob/tot_hjs_glob*100):.1f}% del total</span>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<div style='height: 12px;'></div>", unsafe_allow_html=True)

    # ── 3. Filtros, Ordenamiento y Selector de Vista
    f1, f2, f3, f4, f5 = st.columns([1.8, 1.2, 1.2, 1.2, 1.2])
    with f1:
        q_search = st.text_input("🔍 Búsqueda rápida:", placeholder="Buscar por Orden Interna, PO, OF, Proyecto...", key="search_materia_prima")
    with f2:
        proys_set = set(df_of['proyecto'].dropna().unique())
        if df_po is not None and not df_po.empty:
            proys_set.update(df_po['proyecto'].dropna().unique())
        all_proys = sorted([p for p in proys_set if p and p not in ('Varios', 'General', 'nan', 'POR DEFINIR', '')])
        sel_proy = st.selectbox("Filtrar por Proyecto:", ["Todos"] + all_proys, key="sb_proy_materia_prima")
    with f3:
        sort_by = st.selectbox("Ordenar por:", [
            "🔢 Orden Interna (INT-001 en adelante)",
            "📑 Folio de PO",
            "🏭 Número de OF",
            "📊 Mayor Consumo de Hojas (Descendente)"
        ], index=0, key="sb_sort_materia_prima")
    with f4:
        modo_vista = st.selectbox("Tipo de Reporte:", [
            "📋 Vista por OF (Idéntica a la Imagen)",
            "📦 Vista Consolidada por PO"
        ], key="sb_modo_vista_mp")
    with f5:
        filtro_info = st.selectbox("Información Corte:", [
            "Todas las Órdenes",
            "Con Información de Corte",
            "Sin Información de Corte"
        ], key="sb_filtro_info_mp")

    # Botones y controles
    c_btn1, c_btn2 = st.columns([3, 1.2])
    with c_btn1:
        vista_tabla_tipo = st.radio(
            "Estilo de visualización:",
            ["🎨 Vista de Tabla Formato Oficial (Azul Ejecutivo)", "📊 Vista Interactiva con Filtro por Columnas"],
            horizontal=True,
            key="radio_estilo_tabla_mp"
        )
    with c_btn2:
        # Descarga a Excel
        excel_bytes = generate_materia_prima_excel(df_of, df_po)
        st.download_button(
            label="📥 Descargar a Excel (.xlsx)",
            data=excel_bytes,
            file_name=f"Reporte_Uso_Materia_Prima_{datetime.date.today().strftime('%Y%m%d')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="btn_dl_materia_prima_excel",
            help="Descargar archivo Excel formateado con el diseño azul de cabecera exacto a la imagen solicitada y fórmulas de totales nativas."
        )

    # Filtrar datos
    df_active = df_of.copy() if "Vista por OF" in modo_vista else df_po.copy()
    if q_search:
        q = q_search.strip().lower()
        if "Vista por OF" in modo_vista:
            df_active = df_active[
                df_active['po'].astype(str).str.lower().str.contains(q) |
                df_active['of_number'].astype(str).str.lower().str.contains(q) |
                df_active['proyecto'].astype(str).str.lower().str.contains(q) |
                df_active.get('id_interno', pd.Series(['']*len(df_active))).astype(str).str.lower().str.contains(q)
            ]
        else:
            df_active = df_active[
                df_active['po'].astype(str).str.lower().str.contains(q) |
                df_active['proyecto'].astype(str).str.lower().str.contains(q) |
                df_active.get('id_interno', pd.Series(['']*len(df_active))).astype(str).str.lower().str.contains(q)
            ]
    if sel_proy != "Todos":
        df_active = df_active[df_active['proyecto'] == sel_proy]

    # Filtrar por disponibilidad de información en corte
    if "Con Información" in filtro_info and "informacion" in df_active.columns:
        df_active = df_active[df_active['informacion'] == 'Con Información']
    elif "Sin Información" in filtro_info and "informacion" in df_active.columns:
        df_active = df_active[df_active['informacion'] == 'Sin Información']

    # Aplicar ordenamiento seleccionado
    if "Orden Interna" in sort_by:
        if "of_number" in df_active.columns:
            df_active = df_active.sort_values(by=['int_num', 'po', 'of_number'], ascending=[True, True, True])
        else:
            df_active = df_active.sort_values(by=['int_num', 'po'], ascending=[True, True])
    elif "Folio de PO" in sort_by:
        df_active = df_active.sort_values(by=['po', 'int_num'], ascending=[True, True])
    elif "Número de OF" in sort_by and "of_number" in df_active.columns:
        df_active = df_active.sort_values(by=['of_number'], ascending=True)
    elif "Mayor Consumo" in sort_by:
        df_active = df_active.sort_values(by=['TOTAL HOJAS'], ascending=False)

    # ── 4. Renderizado
    if "Vista de Tabla Formato Oficial" in vista_tabla_tipo:
        if "Vista por OF" in modo_vista:
            # HTML idéntico a la imagen con encabezado azul #2B6CB0 y Orden Interna
            html_code = """
            <!DOCTYPE html>
            <html>
            <head>
            <meta charset="utf-8">
            <style>
                body { margin: 0; padding: 0; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; font-size: 12px; background: transparent; }
                .table-container { overflow-x: auto; max-height: 650px; border: 1px solid #CBD5E1; border-radius: 8px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }
                table { width: 100%; border-collapse: collapse; font-size: 12px; }
                thead th { position: sticky; top: 0; background-color: #2B6CB0; color: #FFFFFF; z-index: 10; padding: 10px 8px; font-weight: 700; text-align: center; border-right: 1px solid #4299E1; }
                thead th:last-child { border-right: none; }
                tbody tr { border-bottom: 1px solid #E2E8F0; background-color: #FFFFFF; }
                tbody tr:nth-child(even) { background-color: #F8FAFC; }
                tbody tr:hover { background-color: #EDF2F7 !important; }
                td { padding: 7px 10px; border-right: 1px solid #E2E8F0; }
                td:last-child { border-right: none; }
                .t-total { background-color: #EDF2F7; font-weight: 800; border-top: 2px solid #0F172A; border-bottom: 2px solid #0F172A; }
            </style>
            </head>
            <body>
            <div class="table-container">
            <table>
                <thead>
                    <tr>
                        <th style="width:105px; text-align:center;">ORDEN<br>INTERNA</th>
                        <th style="width:90px; text-align:center;">PO</th>
                        <th style="min-width:240px; text-align:left;">OF</th>
                        <th style="min-width:110px; text-align:left;">PROYECTO</th>
                        <th style="width:95px; text-align:right;">CAL 10<br>GALV</th>
                        <th style="width:95px; text-align:right;">CAL 12<br>GALV</th>
                        <th style="width:95px; text-align:right;">CAL 14<br>GALV</th>
                        <th style="width:95px; text-align:right;">CAL 16<br>GALV</th>
                        <th style="width:105px; text-align:right;">CAL 10<br>DECAPADO</th>
                        <th style="width:105px; text-align:right;">CAL 12<br>DECAPADO</th>
                        <th style="width:105px; text-align:right;">CAL 14<br>DECAPADO</th>
                        <th style="width:105px; text-align:right;">CAL 16<br>DECAPADO</th>
                        <th style="width:100px; text-align:right; background-color:#1A365D;">TOTAL<br>HOJAS</th>
                    </tr>
                </thead>
                <tbody>
            """
            
            for _, r in df_active.iterrows():
                id_v = str(r.get('id_interno', '')).strip()
                po_v = str(r.get('po', '')).strip()
                of_v = str(r.get('of_number', '')).strip()
                proy_v = str(r.get('proyecto', '')).strip()
                tot_h = float(r.get('TOTAL HOJAS', 0) or 0)
                
                td_mat = ""
                for c_name in STANDARD_COLUMNS:
                    val = float(r.get(c_name, 0) or 0)
                    if val > 0:
                        td_mat += f"<td style='text-align:right; font-weight:bold; color:#2B6CB0;'>{val:,.0f}</td>"
                    else:
                        td_mat += "<td style='text-align:center; color:#CBD5E1;'>-</td>"
                        
                html_code += f"""
                <tr>
                    <td style="text-align:center; font-weight:800; color:#1E3A8A; background-color:#F8FAFC;">{id_v}</td>
                    <td style="text-align:center; font-weight:700; color:#EC2024;">{po_v}</td>
                    <td style="font-weight:600; color:#1E293B;">{of_v}</td>
                    <td style="color:#334155; font-weight:500;">{proy_v}</td>
                    {td_mat}
                    <td style="text-align:right; font-weight:800; color:#0F172A; background-color:#F1F5F9;">{tot_h:,.0f}</td>
                </tr>
                """
                
            # Fila de Totales
            td_tot_mat = ""
            for c_name in STANDARD_COLUMNS:
                s_val = float(df_active[c_name].sum()) if c_name in df_active.columns else 0.0
                td_tot_mat += f"<td style='text-align:right; font-weight:800; color:#0F172A;'>{s_val:,.0f}</td>"
                
            tot_act_h = float(df_active['TOTAL HOJAS'].sum()) if 'TOTAL HOJAS' in df_active.columns else 0.0
            html_code += f"""
                <tr class="t-total">
                    <td colspan="4" style="text-align:center; font-weight:800; color:#0F172A; text-transform:uppercase;">TOTALES GENERALES</td>
                    {td_tot_mat}
                    <td style="text-align:right; font-weight:900; color:#1E3A8A; background-color:#E2E8F0;">{tot_act_h:,.0f}</td>
                </tr>
                </tbody>
            </table>
            </div>
            </body>
            </html>
            """
            calc_h = min(850, max(280, 95 + len(df_active) * 36))
            components.html(html_code, height=calc_h, scrolling=True)
            
        else:
            # Vista Consolidada por PO en Formato Oficial Ejecutivo (Azul)
            html_code = """
            <!DOCTYPE html>
            <html>
            <head>
            <meta charset="utf-8">
            <style>
                body { margin: 0; padding: 0; font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; font-size: 12px; background: transparent; }
                .table-container { overflow-x: auto; max-height: 650px; border: 1px solid #CBD5E1; border-radius: 8px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }
                table { width: 100%; border-collapse: collapse; font-size: 12px; }
                thead th { position: sticky; top: 0; background-color: #2B6CB0; color: #FFFFFF; z-index: 10; padding: 10px 8px; font-weight: 700; text-align: center; border-right: 1px solid #4299E1; }
                thead th:last-child { border-right: none; }
                tbody tr { border-bottom: 1px solid #E2E8F0; background-color: #FFFFFF; }
                tbody tr:nth-child(even) { background-color: #F8FAFC; }
                tbody tr:hover { background-color: #EDF2F7 !important; }
                td { padding: 7px 10px; border-right: 1px solid #E2E8F0; }
                td:last-child { border-right: none; }
                .badge-info { background: #DCFCE7; color: #166534; font-size: 10px; font-weight: 700; padding: 3px 8px; border-radius: 12px; border: 1px solid #BBF7D0; display: inline-block; white-space: nowrap; }
                .badge-no-info { background: #F1F5F9; color: #64748B; font-size: 10px; font-weight: 600; padding: 3px 8px; border-radius: 12px; border: 1px solid #CBD5E1; display: inline-block; white-space: nowrap; }
                .t-total { background-color: #EDF2F7; font-weight: 800; border-top: 2px solid #0F172A; border-bottom: 2px solid #0F172A; }
            </style>
            </head>
            <body>
            <div class="table-container">
            <table>
                <thead>
                    <tr>
                        <th style="width:100px; text-align:center;">ORDEN<br>INTERNA</th>
                        <th style="width:105px; text-align:center;">PO / FOLIO</th>
                        <th style="min-width:140px; text-align:left;">PROYECTO</th>
                        <th style="min-width:140px; text-align:center;">OFs #</th>
                        <th style="width:125px; text-align:center;">INFORMACIÓN</th>
                        <th style="width:90px; text-align:right;">CAL 10<br>GALV</th>
                        <th style="width:90px; text-align:right;">CAL 12<br>GALV</th>
                        <th style="width:90px; text-align:right;">CAL 14<br>GALV</th>
                        <th style="width:90px; text-align:right;">CAL 16<br>GALV</th>
                        <th style="width:95px; text-align:right;">CAL 10<br>DECAPADO</th>
                        <th style="width:95px; text-align:right;">CAL 12<br>DECAPADO</th>
                        <th style="width:95px; text-align:right;">CAL 14<br>DECAPADO</th>
                        <th style="width:95px; text-align:right;">CAL 16<br>DECAPADO</th>
                        <th style="width:95px; text-align:right; background-color:#1A365D;">TOTAL<br>HOJAS</th>
                    </tr>
                </thead>
                <tbody>
            """
            
            for _, r in df_active.iterrows():
                id_v = str(r.get('id_interno', '')).strip()
                po_v = str(r.get('po', '')).strip()
                proy_v = str(r.get('proyecto', '')).strip()
                ofs_v = str(r.get('ofs_count', 'Sin Información')).strip()
                info_v = str(r.get('informacion', 'Sin Información')).strip()
                tot_h = float(r.get('TOTAL HOJAS', 0) or 0)
                
                badge_html = f"<span class='badge-info'>Con Información</span>" if info_v == "Con Información" else f"<span class='badge-no-info'>Sin Información</span>"
                ofs_html = f"<span style='font-weight:700; color:#1E293B;'>{ofs_v}</span>" if ofs_v != "Sin Información" else f"<span style='color:#94A3B8; font-style:italic;'>Sin Información</span>"
                
                td_mat = ""
                for c_name in STANDARD_COLUMNS:
                    val = float(r.get(c_name, 0) or 0)
                    if val > 0:
                        td_mat += f"<td style='text-align:right; font-weight:bold; color:#2B6CB0;'>{val:,.0f}</td>"
                    else:
                        td_mat += "<td style='text-align:center; color:#CBD5E1;'>-</td>"
                        
                tot_str = f"{tot_h:,.0f}" if tot_h > 0 else "<span style='color:#CBD5E1;'>-</span>"
                
                html_code += f"""
                <tr>
                    <td style="text-align:center; font-weight:800; color:#1E3A8A; background-color:#F8FAFC;">{id_v}</td>
                    <td style="text-align:center; font-weight:700; color:#EC2024;">{po_v}</td>
                    <td style="color:#334155; font-weight:500;">{proy_v}</td>
                    <td style="text-align:center;">{ofs_html}</td>
                    <td style="text-align:center;">{badge_html}</td>
                    {td_mat}
                    <td style="text-align:right; font-weight:800; color:#0F172A; background-color:#F1F5F9;">{tot_str}</td>
                </tr>
                """
                
            td_tot_mat = ""
            for c_name in STANDARD_COLUMNS:
                s_val = float(df_active[c_name].sum()) if c_name in df_active.columns else 0.0
                td_tot_mat += f"<td style='text-align:right; font-weight:800; color:#0F172A;'>{s_val:,.0f}</td>"
                
            tot_act_h = float(df_active['TOTAL HOJAS'].sum()) if 'TOTAL HOJAS' in df_active.columns else 0.0
            html_code += f"""
                <tr class="t-total">
                    <td colspan="5" style="text-align:center; font-weight:800; color:#0F172A; text-transform:uppercase;">TOTALES GENERALES CONSOLIDADOS</td>
                    {td_tot_mat}
                    <td style="text-align:right; font-weight:900; color:#1E3A8A; background-color:#E2E8F0;">{tot_act_h:,.0f}</td>
                </tr>
                </tbody>
            </table>
            </div>
            </body>
            </html>
            """
            calc_h = min(850, max(280, 95 + len(df_active) * 36))
            components.html(html_code, height=calc_h, scrolling=True)
            
    else:
        # Modo Interactivo Streamlit Dataframe
        disp_inter = df_active.copy()
        if 'int_num' in disp_inter.columns:
            disp_inter = disp_inter.drop(columns=['int_num'])
        cfg = {
            'id_interno': st.column_config.TextColumn("Orden Interna", help="Folio Interno INT-XXXX"),
            'po': st.column_config.TextColumn("PO / Folio"),
            'of_number': st.column_config.TextColumn("Orden de Fabricación (OF)"),
            'proyecto': st.column_config.TextColumn("Proyecto"),
            'ofs_count': st.column_config.TextColumn("OFs #", help="Números consecutivos de OFs asociadas"),
            'informacion': st.column_config.TextColumn("Información", help="Estatus de información en Corte"),
        }
        for c in STANDARD_COLUMNS + ['TOTAL HOJAS']:
            if c in disp_inter.columns:
                cfg[c] = st.column_config.NumberColumn(c, format="%d hjs" if c == "TOTAL HOJAS" else "%d")
        st.dataframe(
            disp_inter,
            column_config=cfg,
            use_container_width=True,
            hide_index=True
        )

